# Usage:
#   python post_sharepoint_update.py "C:\Users\edvigil\OneDrive\...\SDS_Master_Reconciled_V10_ONEDRIVE.xlsx"
#
# What it does:
#   Finds every row in "📂 SDS File Index" where G = "OUTDATED" and H is not blank,
#   sets F = H (copies portal date to file date), sets G = "CURRENT", and appends
#   a timestamped note to column I.  All other cells, formatting, and formulas are
#   left untouched.  The workbook is saved in-place.

import argparse
import os
import sys
from datetime import date, datetime

import openpyxl

SHEET_NAME = "\U0001F4C2 SDS File Index"  # 📂 SDS File Index

# Column indices (1-based)
COL_DOC    = 1   # A
COL_NAME   = 2   # B  (File Name)
COL_PROD   = 3   # C  (Product Name)
COL_MFR    = 4   # D
COL_CODE   = 5   # E
COL_FDATE  = 6   # F  (Revision Date – our file)
COL_STATUS = 7   # G  (Portal Status)
COL_PDATE  = 8   # H  (Rev Date on Portal)
COL_NOTES  = 9   # I


def _cell_is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def _format_date(value) -> str:
    """Return YYYY-MM-DD string from a date, datetime, or ISO-string value."""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _append_note(existing, addition: str) -> str:
    if _cell_is_blank(existing):
        return addition
    return f"{existing} | {addition}"


def run(workbook_path: str) -> None:
    if not os.path.isfile(workbook_path):
        sys.exit(f"ERROR: file not found: {workbook_path}")

    today_str = date.today().strftime("%Y-%m-%d")

    print(f"Opening: {workbook_path}")
    wb = openpyxl.load_workbook(workbook_path, keep_vba=False)

    if SHEET_NAME not in wb.sheetnames:
        sys.exit(
            f"ERROR: sheet '{SHEET_NAME}' not found.\n"
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb[SHEET_NAME]
    changed_rows = []

    for row in ws.iter_rows(min_row=2):
        g_cell = row[COL_STATUS - 1]
        h_cell = row[COL_PDATE  - 1]

        if g_cell.value != "OUTDATED":
            continue
        if _cell_is_blank(h_cell.value):
            continue

        f_cell    = row[COL_FDATE - 1]
        i_cell    = row[COL_NOTES - 1]
        doc_cell  = row[COL_DOC   - 1]
        prod_cell = row[COL_PROD  - 1]

        old_fdate = _format_date(f_cell.value) if not _cell_is_blank(f_cell.value) else "(blank)"
        new_fdate = _format_date(h_cell.value)

        # ── mutations ─────────────────────────────────────────────────────────
        f_cell.value = h_cell.value                               # F = H
        g_cell.value = "CURRENT"                                  # G = CURRENT
        i_cell.value = _append_note(                              # I += note
            i_cell.value,
            f"Auto-flipped CURRENT {today_str}"
        )

        changed_rows.append({
            "doc":      doc_cell.value  or "",
            "old_date": old_fdate,
            "new_date": new_fdate,
            "product":  prod_cell.value or "",
        })

    if not changed_rows:
        print("No rows matched (G=OUTDATED and H not blank). Nothing changed.")
        return

    wb.save(workbook_path)
    print(f"\nSaved. {len(changed_rows)} row(s) updated.\n")

    # ── summary table ─────────────────────────────────────────────────────────
    col_widths = {
        "doc":      max(6,  max(len(r["doc"])      for r in changed_rows)),
        "old_date": max(12, max(len(r["old_date"]) for r in changed_rows)),
        "new_date": max(12, max(len(r["new_date"]) for r in changed_rows)),
        "product":  max(12, max(len(r["product"])  for r in changed_rows)),
    }
    header = (
        f"{'DOC#':<{col_widths['doc']}}  "
        f"{'Old Rev Date':<{col_widths['old_date']}}  "
        f"{'New Rev Date':<{col_widths['new_date']}}  "
        f"{'Product Name':<{col_widths['product']}}"
    )
    separator = "-" * len(header)
    print(header)
    print(separator)
    for r in changed_rows:
        print(
            f"{r['doc']:<{col_widths['doc']}}  "
            f"{r['old_date']:<{col_widths['old_date']}}  "
            f"{r['new_date']:<{col_widths['new_date']}}  "
            f"{r['product']:<{col_widths['product']}}"
        )
    print(separator)
    print(f"Total updated: {len(changed_rows)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Flip OUTDATED SDS rows to CURRENT using the portal date."
    )
    parser.add_argument("workbook", help="Full path to the .xlsx workbook")
    args = parser.parse_args()
    run(args.workbook)
