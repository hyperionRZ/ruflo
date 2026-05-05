# Usage:
#   python export_vendor_batches.py "C:\Users\edvigil\OneDrive\...\SDS_Master_Reconciled_V10_ONEDRIVE.xlsx"
#
# What it does:
#   Reads all rows in "📂 SDS File Index" where G = "UNVERIFIED" or G is blank.
#   Groups those rows by Manufacturer (column D).
#   For each manufacturer, writes a plain-text file  batch_<Mfr>.txt  whose lines
#   are the K-column pipe-delimited prompt strings.
#   Also writes batch_summary.txt (sorted by row count, descending).
#   All output goes into a  claude_batches/  subfolder beside the workbook.

import argparse
import os
import re
import sys
from collections import defaultdict

import openpyxl

SHEET_NAME = "\U0001F4C2 SDS File Index"  # 📂 SDS File Index

# Column indices (1-based)
COL_DOC    = 1   # A
COL_MFR    = 4   # D  (Manufacturer)
COL_STATUS = 7   # G  (Portal Status)
COL_PROMPT = 11  # K  (Claude Prompt Row)

UNVERIFIED_VALUES = {"UNVERIFIED", "", None}


def _sanitize_filename(name: str, max_len: int = 40) -> str:
    """Replace non-alphanumeric characters with underscores, trim to max_len."""
    if not name or str(name).strip() == "":
        return "UNKNOWN_MANUFACTURER"
    sanitized = re.sub(r"[^A-Za-z0-9]+", "_", str(name).strip())
    sanitized = sanitized.strip("_")
    return sanitized[:max_len] if sanitized else "UNKNOWN_MANUFACTURER"


def run(workbook_path: str) -> None:
    if not os.path.isfile(workbook_path):
        sys.exit(f"ERROR: file not found: {workbook_path}")

    # ── output directory: claude_batches/ next to the workbook ───────────────
    workbook_dir = os.path.dirname(os.path.abspath(workbook_path))
    output_dir   = os.path.join(workbook_dir, "claude_batches")
    os.makedirs(output_dir, exist_ok=True)

    print(f"Opening: {workbook_path}")
    wb = openpyxl.load_workbook(workbook_path, keep_vba=False, read_only=True)

    if SHEET_NAME not in wb.sheetnames:
        sys.exit(
            f"ERROR: sheet '{SHEET_NAME}' not found.\n"
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb[SHEET_NAME]

    # manufacturer -> list of K-column prompt strings
    batches: dict[str, list[str]] = defaultdict(list)

    for row in ws.iter_rows(min_row=2, values_only=True):
        # Pad short rows so index lookups are safe
        row = list(row) + [None] * max(0, COL_PROMPT - len(row))

        g_val = row[COL_STATUS - 1]
        g_str = str(g_val).strip().upper() if g_val is not None else ""

        if g_str not in ("UNVERIFIED", ""):
            continue

        mfr_val    = row[COL_MFR    - 1]
        prompt_val = row[COL_PROMPT - 1]

        mfr_key    = str(mfr_val).strip() if mfr_val is not None else ""
        prompt_str = str(prompt_val).strip() if prompt_val is not None else ""

        if not prompt_str:
            continue  # nothing to export for this row

        batches[mfr_key].append(prompt_str)

    wb.close()

    if not batches:
        print("No UNVERIFIED (or blank) rows with K-column content found.")
        return

    # ── write per-manufacturer files ─────────────────────────────────────────
    summary_rows = []
    for mfr, prompts in batches.items():
        safe_name = _sanitize_filename(mfr)
        filename  = f"batch_{safe_name}.txt"
        filepath  = os.path.join(output_dir, filename)

        with open(filepath, "w", encoding="utf-8") as fh:
            fh.write("\n".join(prompts) + "\n")

        summary_rows.append({
            "manufacturer": mfr or "(blank)",
            "count":        len(prompts),
            "filename":     filename,
        })
        print(f"  Wrote {len(prompts):>4} rows -> {filename}")

    # Sort by count descending
    summary_rows.sort(key=lambda r: r["count"], reverse=True)

    # ── write batch_summary.txt ───────────────────────────────────────────────
    summary_path = os.path.join(output_dir, "batch_summary.txt")
    col_mfr  = max(12, max(len(r["manufacturer"]) for r in summary_rows))
    col_cnt  = 9
    col_file = max(11, max(len(r["filename"])     for r in summary_rows))

    lines = [
        f"{'Manufacturer':<{col_mfr}}  {'Row Count':>{col_cnt}}  {'Output File':<{col_file}}",
        "-" * (col_mfr + col_cnt + col_file + 4),
    ]
    for r in summary_rows:
        lines.append(
            f"{r['manufacturer']:<{col_mfr}}  {r['count']:>{col_cnt}}  {r['filename']:<{col_file}}"
        )
    lines.append("-" * (col_mfr + col_cnt + col_file + 4))
    lines.append(f"Total manufacturers: {len(summary_rows)}")
    lines.append(f"Total rows exported: {sum(r['count'] for r in summary_rows)}")

    with open(summary_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")

    print(f"\nSummary written -> {summary_path}")
    print(f"\n{'Manufacturer':<{col_mfr}}  {'Row Count':>{col_cnt}}  {'Output File':<{col_file}}")
    print("-" * (col_mfr + col_cnt + col_file + 4))
    for r in summary_rows:
        print(f"{r['manufacturer']:<{col_mfr}}  {r['count']:>{col_cnt}}  {r['filename']:<{col_file}}")
    print(f"\nDone. Output folder: {output_dir}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Export UNVERIFIED SDS rows grouped by manufacturer for Claude batch processing."
    )
    parser.add_argument("workbook", help="Full path to the .xlsx workbook")
    args = parser.parse_args()
    run(args.workbook)
